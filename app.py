"""Local HR workbook workspace. Uploaded documents are data, never executable instructions."""
import io, os, json, secrets, sqlite3, threading, zipfile
from pathlib import Path
from datetime import datetime, date, timedelta
from functools import wraps
from flask import Flask, request, session, redirect, render_template, jsonify, send_file, abort
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.workbook.properties import CalcProperties
from docx import Document
from docx.shared import Pt

KINDS = {'transfer':'Điều động', 'termination':'Chấm dứt HĐLĐ', 'title':'Thay đổi chức danh'}

def create_app(data_dir=None):
    app = Flask(__name__)
    root = Path(data_dir or os.environ.get('EMGA_DATA_DIR') or Path(os.environ.get('LOCALAPPDATA', Path.home())) / 'EmgaHR')
    root.mkdir(parents=True, exist_ok=True)
    key = root / 'session.key'
    if not key.exists():
        try:
            with key.open('x') as f: f.write(secrets.token_hex(32))
            key.chmod(0o600)
        except FileExistsError: pass
    app.config.update(SECRET_KEY=key.read_text(), MAX_CONTENT_LENGTH=32*1024*1024,
                      SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Strict', PERMANENT_SESSION_LIFETIME=timedelta(hours=8))
    lock = threading.RLock()
    cache = {}
    def db():
        c = sqlite3.connect(root/'hr.sqlite', timeout=30)
        c.row_factory = sqlite3.Row
        return c
    with db() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL, password TEXT NOT NULL, role TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS books(id TEXT PRIMARY KEY, name TEXT NOT NULL, revision INTEGER NOT NULL DEFAULT 0, created TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS patches(book TEXT, sheet TEXT, cell TEXT, value TEXT NOT NULL, PRIMARY KEY(book,sheet,cell));
        CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, stamp TEXT, actor TEXT, action TEXT, detail TEXT);
        CREATE TABLE IF NOT EXISTS decisions(id INTEGER PRIMARY KEY, stamp TEXT, actor TEXT, payload TEXT);
        ''')
    def audit(c, action, detail):
        c.execute('INSERT INTO audit(stamp,actor,action,detail) VALUES(?,?,?,?)', (datetime.now().isoformat(timespec='seconds'), session.get('user','system'), action, json.dumps(detail, ensure_ascii=False)))
    def auth(role=None):
        def deco(fn):
            @wraps(fn)
            def wrapped(*a, **kw):
                if not session.get('uid'):
                    if request.path.startswith('/api/'): return jsonify(error='Vui lòng đăng nhập.'),401
                    return redirect('/login')
                with db() as c: u=c.execute('SELECT * FROM users WHERE id=?',(session['uid'],)).fetchone()
                if not u: session.clear(); abort(401)
                if role and u['role'] not in role: abort(403)
                return fn(*a, **kw)
            return wrapped
        return deco
    @app.before_request
    def csrf():
        if 'csrf' not in session: session['csrf']=secrets.token_hex(24)
        if request.method=='POST':
            supplied=request.headers.get('X-CSRF-Token') or request.form.get('csrf')
            if not supplied or not secrets.compare_digest(supplied,session['csrf']): abort(403, 'Phiên không hợp lệ; tải lại trang.')
    @app.errorhandler(400)
    @app.errorhandler(403)
    @app.errorhandler(404)
    @app.errorhandler(409)
    @app.errorhandler(413)
    def error(e):
        return jsonify(error=str(e.description)),e.code
    def book(bid):
        with db() as c: b=c.execute('SELECT * FROM books WHERE id=?',(bid,)).fetchone()
        if not b: abort(404,'Không tìm thấy tệp.')
        if bid not in cache:
            cache[bid]=(load_workbook(root/(bid+'.xlsx'), data_only=False),load_workbook(root/(bid+'.xlsx'),data_only=True))
        return b, cache[bid]
    def patches(bid, sheet=None):
        with db() as c:
            rows=c.execute('SELECT * FROM patches WHERE book=?'+(' AND sheet=?' if sheet else ''),(bid,sheet) if sheet else (bid,)).fetchall()
        return {(r['sheet'],r['cell']):json.loads(r['value']) for r in rows}
    def display(v):
        if v is None: return ''
        if isinstance(v,dict) and 'date' in v: return datetime.fromisoformat(v['date']).strftime('%d/%m/%Y')
        if isinstance(v,(date,datetime)): return v.strftime('%d/%m/%Y')
        return str(v)
    @app.route('/login',methods=['GET','POST'])
    def login():
        with db() as c: exists=c.execute('SELECT count(*) FROM users').fetchone()[0]>0
        msg=''
        if request.method=='POST':
            name=request.form.get('name','').strip(); password=request.form.get('password','')
            if not exists:
                # Account bootstrap must originate at the Windows host itself.
                if request.remote_addr not in ('127.0.0.1','::1'): abort(403,'Tạo quản trị viên trên máy chủ trước.')
                if len(password)<12 or not name: msg='Tên không được trống; mật khẩu ít nhất 12 ký tự.'
                else:
                    with lock, db() as c:
                        if c.execute('SELECT count(*) FROM users').fetchone()[0]: abort(409,'Quản trị viên đã được tạo.')
                        c.execute('INSERT INTO users(name,password,role) VALUES(?,?,?)',(name,generate_password_hash(password),'admin'))
                    exists=True
            if exists:
                with db() as c: u=c.execute('SELECT * FROM users WHERE name=?',(name,)).fetchone()
                if u and check_password_hash(u['password'],password):
                    session.clear();session.update(uid=u['id'],user=u['name'],role=u['role'],csrf=secrets.token_hex(24));session.permanent=True
                    return redirect('/')
                msg='Tên đăng nhập hoặc mật khẩu không đúng.'
        return render_template('login.html', setup=not exists,error=msg)
    @app.post('/logout')
    def logout(): session.clear();return redirect('/login')
    @app.get('/')
    @auth()
    def home(): return render_template('index.html', kinds=KINDS)
    @app.get('/api/books')
    @auth()
    def books():
        with db() as c: return jsonify([dict(r) for r in c.execute('SELECT * FROM books ORDER BY created DESC')])
    @app.post('/api/import')
    @auth(('admin','editor'))
    def upload():
        f=request.files.get('file')
        if not f or not f.filename.lower().endswith('.xlsx'): abort(400,'Chỉ nhận tệp .xlsx.')
        raw=f.read()
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist())>150*1024*1024: raise ValueError()
            w=load_workbook(io.BytesIO(raw),data_only=False)
            if sum(s.max_row*s.max_column for s in w)>3_000_000: raise ValueError()
        except Exception: abort(400,'Tệp Excel không hợp lệ hoặc vượt giới hạn dữ liệu.')
        bid=secrets.token_hex(16)
        with lock:
            (root/(bid+'.xlsx')).write_bytes(raw)
            with db() as c:
                c.execute('INSERT INTO books(id,name,created) VALUES(?,?,?)',(bid,Path(f.filename).name,datetime.now().isoformat(timespec='seconds')))
                audit(c,'import',{'book':bid,'name':Path(f.filename).name})
        return jsonify(id=bid)
    @app.get('/api/book/<bid>')
    @auth()
    def metadata(bid):
        with lock:
            b,(w,_) = book(bid)
            return jsonify(**dict(b), external_links=len(w._external_links),sheets=[dict(name=s.title,rows=s.max_row,cols=s.max_column) for s in w])
    @app.get('/api/book/<bid>/sheet')
    @auth()
    def sheet(bid):
        with lock:
            b,(w,v)=book(bid); name=request.args.get('sheet',w.sheetnames[0])
            if name not in w.sheetnames: abort(404,'Không có trang tính.')
            s=w[name]; cached=v[name]; p=patches(bid,name)
            try:
                start=max(1,int(request.args.get('row',1))); col=max(1,int(request.args.get('col',1)))
            except ValueError: abort(400,'Vị trí không hợp lệ.')
            q=request.args.get('q','').strip().casefold()
            rows=range(start,min(start+30,s.max_row+1))
            if q:
                rows=[]
                for r in range(start,s.max_row+1):
                    # Search raw literals and Excel's last saved results; no formulas are executed.
                    if any(q in display(p.get((name,c.coordinate),cached.cell(r,c.column).value if c.data_type=='f' else c.value)).casefold() for c in s[r]): rows.append(r)
                    if len(rows)==30: break
            out=[]
            for r in rows:
                cells=[]
                for cc in range(col,min(col+16,s.max_column+1)):
                    cell=s.cell(r,cc); formula=cell.data_type=='f'; coord=cell.coordinate
                    value=p.get((name,coord),cached.cell(r,cc).value if formula else cell.value)
                    cells.append(dict(cell=coord,value=display(value),formula=formula,expression=cell.value if formula else '',merged=isinstance(cell,MergedCell),edited=(name,coord) in p, date=isinstance(cell.value,(date,datetime))))
                out.append(dict(row=r,cells=cells))
            return jsonify(revision=b['revision'],rows=out,columns=[get_column_letter(c) for c in range(col,min(col+16,s.max_column+1))], total_rows=s.max_row,total_cols=s.max_column)
    @app.post('/api/book/<bid>/cell')
    @auth(('admin','editor'))
    def edit(bid):
        data=request.get_json() or {}
        with lock:
            b,(w,_) = book(bid)
            try:
                name=data['sheet'];coord=data['cell'];s=w[name];cell=s[coord]
                if not hasattr(cell,'coordinate') or cell.row>s.max_row or cell.column>s.max_column: raise ValueError()
                if len(str(data.get('value','')))>10000: raise ValueError()
            except (KeyError,ValueError,TypeError): abort(400,'Ô không hợp lệ.')
            if cell.data_type=='f' or isinstance(cell,MergedCell): abort(400,'Ô công thức hoặc ô gộp không thể sửa trực tiếp.')
            value=data.get('value',''); typ=data.get('type','text')
            if typ=='number':
                try:
                    value=float(value)
                    if not __import__('math').isfinite(value): raise ValueError()
                except (ValueError,TypeError): abort(400,'Số không hợp lệ.')
            elif typ=='date':
                try: value={'date':datetime.strptime(value,'%Y-%m-%d').isoformat()}
                except (ValueError,TypeError): abort(400,'Ngày phải có dạng yyyy-mm-dd.')
            elif typ=='empty': value=None
            elif typ!='text' or not isinstance(value,str): abort(400,'Kiểu dữ liệu không hợp lệ.')
            if isinstance(value,str) and value.startswith(('=','+','-','@')): abort(400,'Không nhập công thức vào ô văn bản; dùng kiểu Số cho số âm.')
            with db() as c:
                if data.get('revision')!=b['revision']: abort(409,'Người khác vừa sửa tệp. Tải lại bảng và thử lại.')
                old=patches(bid,name).get((name,coord),display(cell.value))
                c.execute('INSERT OR REPLACE INTO patches VALUES(?,?,?,?)',(bid,name,coord,json.dumps(value,ensure_ascii=False)))
                c.execute('UPDATE books SET revision=revision+1 WHERE id=?',(bid,))
                audit(c,'edit',dict(book=bid,sheet=name,cell=coord,old=old,new=value))
            return jsonify(ok=True,revision=b['revision']+1)
    @app.get('/api/book/<bid>/export')
    @auth()
    def export(bid):
        with lock:
            b,_=book(bid)
            w=load_workbook(root/(bid+'.xlsx'))
            for (sheet,coord),value in patches(bid).items():
                if isinstance(value,dict) and 'date' in value: value=datetime.fromisoformat(value['date'])
                w[sheet][coord]=value
            w.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True,forceFullCalc=True)
            result=io.BytesIO();w.save(result);result.seek(0)
        return send_file(result,as_attachment=True,download_name='Bao_cao_cap_nhat.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    @app.get('/api/book/<bid>/people')
    @auth()
    def people(bid):
        with lock:
            _,(w,_) = book(bid)
            if 'DU LIEU CO SO' not in w.sheetnames: return jsonify([])
            s=w['DU LIEU CO SO'];p=patches(bid,s.title);out=[]
            for r in range(4,s.max_row+1):
                def val(c):
                    x=s.cell(r,c);return display(p.get((s.title,x.coordinate),x.value if x.data_type!='f' else ''))
                if val(6):out.append(dict(row=r,code=val(5),name=val(6),title=val(2),unit=val(14)))
            return jsonify(out)
    @app.get('/api/users')
    @auth(('admin',))
    def users():
        with db() as c: return jsonify([dict(r) for r in c.execute('SELECT id,name,role FROM users')])
    @app.post('/api/users')
    @auth(('admin',))
    def add_user():
        d=request.get_json() or {}; name=str(d.get('name','')).strip();pw=d.get('password','');role=d.get('role')
        if not name or len(name)>80 or not isinstance(pw,str) or len(pw)<12 or role not in ('admin','editor','viewer'): abort(400,'Tên hợp lệ, mật khẩu từ 12 ký tự và vai trò bắt buộc.')
        try:
            with db() as c:
                c.execute('INSERT INTO users(name,password,role) VALUES(?,?,?)',(name,generate_password_hash(pw),role));audit(c,'user_created',dict(name=name,role=role))
        except sqlite3.IntegrityError: abort(400,'Tên đăng nhập đã tồn tại.')
        return jsonify(ok=True)
    @app.get('/api/audit')
    @auth(('admin',))
    def history():
        with db() as c: return jsonify([dict(r) for r in c.execute('SELECT * FROM audit ORDER BY id DESC LIMIT 200')])
    @app.get('/api/backup')
    @auth(('admin',))
    def backup():
        out=io.BytesIO()
        with lock, zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
            temp=sqlite3.connect(':memory:')
            with db() as c:c.backup(temp)
            z.writestr('hr.sqlite',temp.serialize());temp.close()
            for p in root.glob('*.xlsx'): z.write(p,p.name)
            for p in root.glob('template-*.docx'): z.write(p,p.name)
        out.seek(0)
        return send_file(out,as_attachment=True,download_name='Emga_sao_luu_'+date.today().isoformat()+'.zip')
    @app.post('/api/templates/<kind>')
    @auth(('admin',))
    def save_template(kind):
        if kind not in KINDS: abort(404)
        f=request.files.get('file')
        if not f or not f.filename.lower().endswith('.docx'): abort(400,'Cần tệp Word .docx.')
        raw=f.read()
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist())>30*1024*1024:raise ValueError()
            Document(io.BytesIO(raw))
        except Exception: abort(400,'Tệp Word không hợp lệ.')
        with lock:
            (root/('template-'+kind+'.docx')).write_bytes(raw)
            with db() as c:audit(c,'template_uploaded',dict(kind=kind))
        return jsonify(ok=True)
    @app.post('/api/decisions')
    @auth(('admin','editor'))
    def decision():
        d=request.get_json() or {};kind=d.get('kind')
        if kind not in KINDS: abort(400,'Loại quyết định không hợp lệ.')
        for k in ['number','effective','reason']:
            if not isinstance(d.get(k),str) or not d[k].strip(): abort(400,'Điền số quyết định, ngày hiệu lực và nội dung/căn cứ.')
        try:datetime.strptime(d['effective'],'%Y-%m-%d')
        except ValueError:abort(400,'Ngày hiệu lực không hợp lệ.')
        with lock:
            _,(w,_) = book(d.get('book'));s=w['DU LIEU CO SO'] if 'DU LIEU CO SO' in w.sheetnames else None
            try:r=int(d['row']);assert s is not None and 4<=r<=s.max_row
            except (KeyError,ValueError,TypeError,AssertionError):abort(400,'Chọn nhân sự từ dữ liệu cơ sở.')
            p=patches(d['book'],s.title)
            def val(c):
                x=s.cell(r,c);return display(p.get((s.title,x.coordinate),x.value if x.data_type!='f' else ''))
            fields={ 'ho_ten':val(6),'ma_nv':val(5),'chuc_danh_cu':val(2),'don_vi_cu':val(14),
                    'so_quyet_dinh':d['number'],'ngay_hieu_luc':datetime.strptime(d['effective'],'%Y-%m-%d').strftime('%d/%m/%Y'),
                    'chuc_danh_moi':str(d.get('new_title','')),'don_vi_moi':str(d.get('new_unit','')),
                    'can_cu':d['reason'],'nguoi_ky':str(d.get('signer','')),'loai_quyet_dinh':KINDS[kind]}
            if not fields['ho_ten']:abort(400,'Dòng không có họ tên.')
            template=root/('template-'+kind+'.docx')
            if template.exists():
                doc=Document(template)
                def replace_paras(paras):
                    for para in paras:
                        # Replace across Word runs without flattening unrelated text or styling.
                        for k,v in fields.items():
                            marker='{{'+k+'}}'
                            while marker in para.text:
                                start=para.text.index(marker);end=start+len(marker);offset=0;hits=[]
                                for run in para.runs:
                                    a=offset;b=a+len(run.text);offset=b
                                    if b>start and a<end:hits.append((run,a,b))
                                if not hits:break
                                first,a,b=hits[0];prefix=first.text[:start-a]
                                if len(hits)==1:first.text=prefix+v+first.text[end-a:]
                                else:
                                    suffix=hits[-1][0].text[end-hits[-1][1]:]
                                    first.text=prefix+v
                                    for run,_,_ in hits[1:-1]:run.text=''
                                    hits[-1][0].text=suffix
                def container(obj):
                    replace_paras(obj.paragraphs)
                    for table in obj.tables:
                        for row in table.rows:
                            for cell in row.cells:container(cell)
                container(doc)
                for section in doc.sections:container(section.header);container(section.footer)
            else:
                doc=Document();doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(13)
                doc.add_paragraph('BẢN DỰ THẢO — CHƯA ÁP DỤNG MẪU ĐƠN VỊ',style='Subtitle')
                doc.add_heading('QUYẾT ĐỊNH '+KINDS[kind].upper(),0)
                doc.add_paragraph('Số: '+fields['so_quyet_dinh'])
                doc.add_paragraph('Căn cứ / nội dung do người lập cung cấp:\n'+fields['can_cu'])
                doc.add_paragraph('Người lao động: '+fields['ho_ten']+' — Mã: '+fields['ma_nv'])
                doc.add_paragraph('Chức danh hiện tại: '+fields['chuc_danh_cu']+'; Đơn vị: '+fields['don_vi_cu'])
                if kind!='termination':doc.add_paragraph('Chức danh mới: '+fields['chuc_danh_moi']+'; Đơn vị mới: '+fields['don_vi_moi'])
                doc.add_paragraph('Ngày hiệu lực dự kiến: '+fields['ngay_hieu_luc'])
                doc.add_paragraph('Người ký: '+fields['nguoi_ky'])
                doc.add_paragraph('Cần đối chiếu căn cứ pháp lý, thẩm quyền, quyền lợi và mẫu đơn vị trước khi ký. Bản này không tự cập nhật hồ sơ hay chấm dứt hợp đồng.')
            out=io.BytesIO();doc.save(out);out.seek(0)
            with db() as c:
                c.execute('INSERT INTO decisions(stamp,actor,payload) VALUES(?,?,?)',(datetime.now().isoformat(),session['user'],json.dumps(fields,ensure_ascii=False)))
                audit(c,'decision_created',dict(kind=kind,book=d['book'],row=r,number=d['number'],custom_template=template.exists()))
        return send_file(out,as_attachment=True,download_name='Du_thao_quyet_dinh.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    from hr import register_hr
    register_hr(app, db, lock, auth, audit, book, patches, root)
    return app

if __name__=='__main__':
    from waitress import serve
    port=int(os.environ.get('EMGA_PORT','8765'))
    print(f'Emga HR: port {port}. Open the browser on the host to create the first administrator.',flush=True)
    serve(create_app(),host=os.environ.get('EMGA_HOST','127.0.0.1'),port=port,threads=4)

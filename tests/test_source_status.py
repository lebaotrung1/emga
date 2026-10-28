import io,unittest
from datetime import datetime
from openpyxl import Workbook,load_workbook
import test_app

class SourceStatus(unittest.TestCase):
    setUp=test_app.Workflow.setUp
    tearDown=test_app.Workflow.tearDown
    def post(self,path,data):
        r=self.client.post(path,json=data,headers=self.h);self.assertEqual(r.status_code,200,r.get_data(as_text=True)[:500]);return r.json
    def upload(self,rows):
        w=Workbook();s=w.active;s.title='DU LIEU CO SO';s['GO2']='Cuối kì';s['HE2']='Cuối kỳ'
        for i,item in enumerate(rows,4):
            s.cell(i,5,item.get('code','T'+str(i)));s.cell(i,6,item.get('name','Nhân sự thử '+str(i)));s.cell(i,7,datetime(1990,1,1));s.cell(i,13,datetime(2020,1,1));s.cell(i,2,'Kỹ thuật viên');s.cell(i,206,item.get('unit','PX1'));s.cell(i,51,'ĐH')
            if 'end' in item:s.cell(i,197,item['end']);s.cell(i,213,item.get('tkv',item['end']))
        raw=io.BytesIO();w.save(raw);raw.seek(0)
        r=self.client.post('/api/import',headers=self.h,data={'file':(raw,'closing-codes.xlsx')});self.assertEqual(r.status_code,200);return r.json['id']
    def test_import_status_by_closing_code_and_keep_retired(self):
        bid=self.upload([dict(code='ACTIVE',end='PX1CÔ'),dict(code='LEFT',end=None),dict(code='ERROR',end='#REF!')])
        preview=self.client.get('/api/hr/import/'+bid).json
        self.assertEqual([p['status'] for p in preview],['active','terminated','review'])
        self.post('/api/hr/import/'+bid,dict(revision=0,effective='2026-09-30'))
        all_people=self.client.get('/api/hr/people?asof=2026-09-30').json;self.assertEqual(len(all_people),3)
        by_code={p['code']:p for p in all_people};self.assertEqual(by_code['LEFT']['status'],'terminated');self.assertEqual(by_code['LEFT']['source_evidence']['end_column'],'GO')
        retired=self.client.get('/api/hr/people?asof=2026-09-30&status=terminated').json;self.assertEqual([p['code'] for p in retired],['LEFT'])
        report=self.client.get('/api/hr/reports?start=2026-09-01&end=2026-09-30').json
        self.assertEqual(report['total'],1);self.assertEqual(report['review_count'],1);self.assertEqual(len(report['people']),3)
        self.assertTrue(self.client.get('/api/hr/people/'+by_code['LEFT']['id']+'/history').json)
    def test_source_filter_and_filtered_export_preserve_original(self):
        bid=self.upload([dict(code='ACTIVE',end='PX1CÔ'),dict(code='LEFT',end=0,unit='PX2'),dict(code='CONFLICT',end='PX1CÔ',tkv=None)])
        meta=self.client.get('/api/book/'+bid).json
        self.assertEqual(meta['source_info']['status_counts'],dict(active=1,terminated=1,review=1))
        base='/api/book/'+bid
        r=self.client.get(base+'/sheet',query_string=dict(sheet='DU LIEU CO SO',status='terminated',unit='PX2',education='university'))
        self.assertEqual(r.status_code,200);self.assertEqual([x['row'] for x in r.json['rows']],[5]);self.assertEqual(r.json['rows'][0]['source_status']['status'],'terminated')
        raw=self.client.get(base+'/filtered-export?status=terminated&unit=PX2').data;w=load_workbook(io.BytesIO(raw))
        self.assertEqual(w.active.max_row,2);self.assertEqual(w.active['B2'].value,'LEFT');self.assertEqual(w.active['J2'].value,'Đã nghỉ')
        raw=self.client.get(base+'/export').data;w=load_workbook(io.BytesIO(raw));self.assertEqual(w.active['E4'].value,'ACTIVE');self.assertEqual(w.active['E5'].value,'LEFT');self.assertEqual(w.active['E6'].value,'CONFLICT')
    def test_error_or_missing_formula_result_requires_review(self):
        bid=self.upload([dict(end='=1+1'),dict(end='PX1CÔ',tkv=None),dict(end=''),dict(end='0')])
        self.assertEqual([p['status'] for p in self.client.get('/api/hr/import/'+bid).json],['review','review','terminated','terminated'])
    def test_source_filter_pagination_has_no_duplicates_or_skips(self):
        bid=self.upload([dict(code='P'+str(i),end='PX1CÔ' if i%2==0 else None) for i in range(70)])
        url='/api/book/'+bid+'/sheet';q=dict(sheet='DU LIEU CO SO',status='active',row=1)
        first=self.client.get(url,query_string=q).json;self.assertEqual(len(first['rows']),30)
        q['row']=first['next_row'];second=self.client.get(url,query_string=q).json
        rows=[x['row'] for x in first['rows']+second['rows']];self.assertEqual(rows,list(range(4,74,2)));self.assertIsNone(second['next_row'])
    def test_sync_old_review_status_is_audited_and_preserves_history(self):
        bid=self.upload([dict(code='ACTIVE',end='PX1CÔ'),dict(code='LEFT',end=None)])
        self.post('/api/hr/import/'+bid,dict(revision=0,effective='2026-09-30'))
        old=self.client.get('/api/hr/people?asof=2026-09-30').json
        # Simulate a previous-version review or a mistaken status confirmation.
        for p in old:
            p.update(effective='2026-10-01',status='review');self.post('/api/hr/people/'+p['id'],p)
        changes=self.client.get('/api/hr/import/'+bid+'/sync-status').json;self.assertEqual(len(changes),2)
        payload=dict(revision=0,effective='2026-10-02',selected=[dict(id=p['id'],version=p['version']) for p in changes])
        self.assertEqual(self.post('/api/hr/import/'+bid+'/sync-status',payload)['updated'],2)
        then=self.client.get('/api/hr/people?asof=2026-10-01').json;self.assertTrue(all(p['status']=='review' for p in then))
        now=self.client.get('/api/hr/people?asof=2026-10-02').json;self.assertEqual({p['code']:p['status'] for p in now},dict(ACTIVE='active',LEFT='terminated'))
        r=self.client.post('/api/hr/import/'+bid+'/sync-status',json=payload,headers=self.h);self.assertEqual(r.status_code,409)
        self.assertEqual(len(self.client.get('/api/hr/people?asof=2026-10-02').json),2)
    def test_retired_person_without_employee_code_is_archived(self):
        bid=self.upload([dict(code='',name='Người đã nghỉ thiếu mã',end=None),dict(code='',name='Người đang làm thiếu mã',end='PX1CÔ')])
        preview=self.client.get('/api/hr/import/'+bid).json
        self.assertTrue(all(p['eligible'] for p in preview))
        self.assertEqual(len({p['storage_code'] for p in preview}),2)
        self.post('/api/hr/import/'+bid,dict(revision=0,effective='2026-09-30'))
        people=self.client.get('/api/hr/people?asof=2026-09-30').json
        self.assertEqual(len(people),2);self.assertEqual({p['status'] for p in people},{'active','terminated'})
        self.assertTrue(all(p['source_evidence']['source_code']=='' for p in people))
        self.assertEqual(self.post('/api/hr/import/'+bid,dict(revision=0,effective='2026-09-30'))['imported'],0)
    def test_duplicate_employee_code_keeps_source_rows_and_closing_status(self):
        bid=self.upload([dict(code='REUSED',name='Dòng thứ nhất thử',end='PX1CÔ'),dict(code='REUSED',name='Dòng thứ hai thử',end=None)])
        preview=self.client.get('/api/hr/import/'+bid).json
        self.assertEqual([p['status'] for p in preview],['active','terminated'])
        self.assertTrue(all(p['storage_code'].startswith('LUU-TRU-') for p in preview))
        self.post('/api/hr/import/'+bid,dict(revision=0,effective='2026-09-30'))
        people=self.client.get('/api/hr/people?asof=2026-09-30').json
        self.assertEqual(len(people),2);self.assertEqual(sum(p['status']=='active' for p in people),1)
        self.assertTrue(all(any('trùng' in w for w in p['warnings']) for p in people))
    def test_normalized_personnel_filters_combine(self):
        bid=self.upload([dict(code='ACTIVE',end='PX1CÔ'),dict(code='LEFT',end=None,unit='PX2')]);self.post('/api/hr/import/'+bid,dict(revision=0,effective='2026-09-30'))
        path='/api/hr/people?asof=2026-09-30&status=active&unit=PX1&education=university&group=unknown&age_group=36%E2%80%9345'
        self.assertEqual([p['code'] for p in self.client.get(path).json],['ACTIVE'])
        self.assertEqual(self.client.get(path+'&q=khong-co').json,[])

if __name__=='__main__':unittest.main()

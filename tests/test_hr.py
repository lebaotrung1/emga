import io, unittest
from openpyxl import load_workbook
import test_app
from hr import age_group

class PersonnelWorkflow(unittest.TestCase):
    setUp=test_app.Workflow.setUp
    tearDown=test_app.Workflow.tearDown
    def post(self,url,data,status=200):
        r=self.client.post(url,json=data,headers=self.h);self.assertEqual(r.status_code,status,r.get_data(as_text=True)[:500]);return r
    def job(self,title,unit,group,technical=False,grade=None,effective='2020-01-01'):
        jid=self.post('/api/hr/jobs',dict(title=title,unit=unit)).json['id']
        self.post('/api/hr/jobs/'+jid+'/rules',dict(version=0,effective=effective,labor_group=group,technical=technical,grade=grade));return jid
    def hire(self,code,jid,birth='2001-11-01',education='university'):
        return self.post('/api/hr/people',dict(code=code,name='Nhân sự '+code,birth_date=birth,hire_date='2026-01-01',gender='female',education=education,job_id=jid,status='active',effective='2026-01-01')).json['id']
    def people(self,asof):return self.client.get('/api/hr/people?asof='+asof).json
    def report(self,start,end):
        r=self.client.get('/api/hr/reports',query_string=dict(start=start,end=end));self.assertEqual(r.status_code,200);return r.json
    def draft(self,pid,jid,kind='transfer',effective='2026-11-01',version=1):
        return self.post('/api/hr/decisions',dict(person=pid,version=version,kind=kind,new_job=jid,effective=effective,number='QD-01',reason='Nội dung thử nghiệm',signer='Người ký thử')).json['id']
    def test_age_boundaries_and_birthday(self):
        self.assertEqual(age_group('2001-11-01','2026-10-31'),(24,'Dưới 25'))
        self.assertEqual(age_group('2001-11-01','2026-11-01'),(25,'25–35'))
        for age,expected in [(35,'25–35'),(36,'36–45'),(45,'36–45'),(46,'46–55'),(55,'46–55'),(56,'56–60'),(60,'56–60'),(61,'Trên 60')]:
            self.assertEqual(age_group(str(2026-age)+'-01-01','2026-01-01')[1],expected)
        self.assertEqual(age_group('2000-02-29','2025-02-28')[0],24)
        self.assertEqual(age_group('2000-02-29','2025-03-01')[0],25)
        self.assertEqual(age_group('','2026-01-01'),(None,'Chưa xác định'))
        self.assertEqual(age_group('1985-01-01','2026-01-01','primary')[1],'41–50')
    def test_hire_auto_classification_and_report_export(self):
        jid=self.job('Công nhân bậc 3','PX1','technology',True,3);pid=self.hire('NV01',jid,education='vocational')
        p=self.people('2026-10-31')[0];self.assertEqual(p['age_group'],'Dưới 25');self.assertEqual(p['job']['labor_group'],'technology');self.assertEqual(p['job']['grade'],3)
        d=self.report('2026-01-01','2026-10-31');self.assertEqual((d['total'],d['female'],d['technical']),(1,1,1));self.assertEqual(d['education'],{'CNKT':1})
        self.assertEqual(d['units'][0],dict(unit='PX1',opening=0,increase=1,decrease=0,closing=1))
        r=self.client.get('/api/hr/reports/export?start=2026-01-01&end=2026-10-31');self.assertEqual(r.status_code,200)
        w=load_workbook(io.BytesIO(r.data));self.assertEqual(len(w.sheetnames),4);self.assertEqual(w['Danh sách chốt kỳ']['A2'].value,'NV01')
    def test_future_transfer_apply_and_historical_reports(self):
        old=self.job('Công nhân','PX1','technology',True,3);new=self.job('Quản đốc','PX2','management')
        pid=self.hire('NV01',old);did=self.draft(pid,new)
        self.assertEqual(self.people('2026-12-01')[0]['job']['unit'],'PX1')
        self.assertEqual(self.client.get('/api/hr/decisions/'+did+'/word').status_code,200)
        self.post('/api/hr/decisions/'+did+'/apply',{})
        self.post('/api/hr/decisions/'+did+'/apply',{})
        self.assertEqual(self.people('2026-10-31')[0]['job']['unit'],'PX1')
        p=self.people('2026-11-01')[0];self.assertEqual(p['job']['labor_group'],'management');self.assertEqual(p['age_group'],'25–35');self.assertEqual(p['education'],'university')
        d=self.report('2026-11-01','2026-11-30');self.assertEqual(d['total'],1);self.assertEqual(d['groups'],{'Quản lý':1})
        self.assertEqual(d['units'],[dict(unit='PX1',opening=1,increase=0,decrease=1,closing=0),dict(unit='PX2',opening=0,increase=1,decrease=0,closing=1)])
        history=self.client.get('/api/hr/people/'+pid+'/history').json;self.assertEqual(len(history),2)
    def test_termination_and_concurrent_decision_protection(self):
        jid=self.job('Kỹ thuật viên','PX1','technology');pid=self.hire('NV01',jid)
        did=self.draft(pid,jid,'termination');other=self.draft(pid,jid,'title')
        self.post('/api/hr/decisions/'+did+'/apply',{})
        self.post('/api/hr/decisions/'+other+'/apply',{},409)
        self.assertEqual(self.report('2026-11-01','2026-11-30')['total'],0)
        self.assertEqual(self.report('2026-10-01','2026-10-31')['total'],1)
    def test_rule_and_education_effective_dates(self):
        jid=self.job('Kỹ thuật viên','PX1','technology');pid=self.hire('NV01',jid)
        self.post('/api/hr/jobs/'+jid+'/rules',dict(version=1,effective='2026-09-01',labor_group='auxiliary',technical=False,grade=None))
        self.assertEqual(self.people('2026-08-31')[0]['job']['labor_group'],'technology')
        self.assertEqual(self.people('2026-09-01')[0]['job']['labor_group'],'auxiliary')
        p=self.people('2026-10-01')[0];p.update(effective='2026-10-01',education='postgraduate')
        self.post('/api/hr/people/'+pid,p)
        self.assertEqual(self.people('2026-09-30')[0]['education'],'university');self.assertEqual(self.people('2026-10-01')[0]['education'],'postgraduate')
        self.post('/api/hr/people/'+pid,p,409)
        p.update(version=2,effective='2026-08-01');self.post('/api/hr/people/'+pid,p,409)
    def test_import_review_dedup_and_baseline_warning(self):
        r=self.client.get('/api/hr/import/'+self.bid);self.assertEqual(r.status_code,200);self.assertEqual(len(r.json),1)
        d={'effective':'2026-10-01','revision':0}
        self.assertEqual(self.post('/api/hr/import/'+self.bid,d).json['imported'],1)
        self.assertEqual(self.post('/api/hr/import/'+self.bid,d).json['imported'],0)
        p=self.people('2026-10-01')[0];self.assertEqual(p['status'],'review');self.assertEqual(p['birth_date'],'1990-02-01')
        report=self.report('2026-01-01','2026-10-31');self.assertEqual(report['total'],0);self.assertEqual(report['review_count'],1);self.assertTrue(report['warnings'])
    def test_validation_and_duplicate_codes(self):
        jid=self.job('Kỹ thuật viên','PX1','technology');self.hire('NV01',jid)
        self.post('/api/hr/people',dict(code='NV01',effective='2026-01-01'),409)
        self.post('/api/hr/people',dict(code='NV02',name='Thử',effective='2026-01-01',status='active',education='university',gender='',birth_date='2030-01-01',hire_date='',job_id=jid),400)
        self.post('/api/hr/jobs/'+jid+'/rules',dict(version=1,effective='2020-01-01',labor_group='service',technical=False,grade=None),409)
    def test_import_does_not_pick_arbitrary_duplicate(self):
        raw=self.client.get('/api/book/'+self.bid+'/export').data
        w=load_workbook(io.BytesIO(raw));s=w['DU LIEU CO SO'];s['E5']='001';s['F5']='Người trùng mã kiểm thử'
        f=io.BytesIO();w.save(f);f.seek(0)
        r=self.client.post('/api/import',headers=self.h,data={'file':(f,'duplicates.xlsx')});self.assertEqual(r.status_code,200);bid=r.json['id']
        preview=self.client.get('/api/hr/import/'+bid).json
        self.assertEqual(len(preview),2);self.assertTrue(all(not p['eligible'] for p in preview))
        r=self.post('/api/hr/import/'+bid,dict(effective='2026-10-01',revision=0));self.assertEqual(r.json['imported'],0);self.assertEqual(r.json['skipped'],2)
    def test_batch_atomic_updates_and_history(self):
        jid=self.job('Kỹ thuật viên','PX1','technology');p1=self.hire('NV01',jid);p2=self.hire('NV02',jid)
        data=dict(effective='2026-10-01',selected=[dict(id=p1,version=1),dict(id=p2,version=99)],changes=dict(education='postgraduate'))
        self.post('/api/hr/people/batch',data,409)
        self.assertTrue(all(p['education']=='university' for p in self.people('2026-10-01')))
        data['selected'][1]['version']=1
        self.assertEqual(self.post('/api/hr/people/batch',data).json['updated'],2)
        self.assertTrue(all(p['education']=='postgraduate' for p in self.people('2026-10-01')))
        self.assertTrue(all(p['education']=='university' for p in self.people('2026-09-30')))
    def test_backup_restore_normalized_data_and_existing_users(self):
        import tempfile,zipfile
        from pathlib import Path
        from app import create_app
        jid=self.job('Kỹ thuật viên','PX1','technology');pid=self.hire('NV01',jid)
        raw=self.client.get('/api/backup').data
        with tempfile.TemporaryDirectory() as dest:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:z.extractall(dest)
            app=create_app(dest);client=app.test_client();client.get('/login')
            with client.session_transaction() as sess:token=sess['csrf']
            self.assertEqual(client.post('/login',data=dict(csrf=token,name='admin',password='a long passphrase 123')).status_code,302)
            r=client.get('/api/hr/people?asof=2026-10-01');self.assertEqual(r.status_code,200);self.assertEqual(r.json[0]['id'],pid)
            self.assertEqual(r.json[0]['job']['labor_group'],'technology')
    def test_report_explains_missing_data_review_and_wrong_period(self):
        r=self.report('2026-01-01','2026-10-31')
        self.assertEqual(r['readiness']['state'],'no_people');self.assertEqual(r['readiness']['source_books'],1)
        self.assertIn('chưa nhập thành hồ sơ',r['readiness']['message'])
        self.post('/api/hr/import/'+self.bid,dict(effective='2026-10-01',revision=0))
        r=self.report('2026-07-01','2026-09-30')
        self.assertEqual(r['readiness']['state'],'before_records');self.assertEqual(r['readiness']['first_effective'],'2026-10-01')
        r=self.report('2026-10-01','2026-10-31')
        self.assertEqual(r['readiness']['state'],'needs_review');self.assertEqual(r['readiness']['review_count'],1)
        raw=self.client.get('/api/hr/reports/export?start=2026-10-01&end=2026-10-31').data
        w=load_workbook(io.BytesIO(raw));values=[str(c.value or '') for row in w.worksheets[0] for c in row]
        self.assertIn('Tình trạng dữ liệu',values);self.assertTrue(any('chờ xác nhận' in v for v in values))
    def test_report_failure_is_not_a_valid_zero_report(self):
        r=self.client.get('/api/hr/reports?start=2026-11-01&end=2026-10-01')
        self.assertEqual(r.status_code,400);self.assertIn('error',r.json);self.assertNotIn('total',r.json)
    def test_viewer_cannot_mutate_hr(self):
        self.post('/api/users',dict(name='reader',password='reader password 123',role='viewer'))
        client=self.app.test_client();client.get('/login')
        with client.session_transaction() as s:t=s['csrf']
        client.post('/login',data=dict(csrf=t,name='reader',password='reader password 123'))
        with client.session_transaction() as s:t=s['csrf']
        for url in ['/api/hr/people','/api/hr/jobs','/api/hr/decisions','/api/hr/jobs/test/rules','/api/hr/decisions/test/apply','/api/hr/import/'+self.bid]:
            self.assertEqual(client.post(url,json={},headers={'X-CSRF-Token':t}).status_code,403,url)
        self.assertEqual(client.get('/api/hr/reports').status_code,200)

if __name__=='__main__':unittest.main()

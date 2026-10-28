import io, tempfile, unittest, re, zipfile
from openpyxl import Workbook, load_workbook
from docx import Document
from app import create_app

class Workflow(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.app=create_app(self.tmp.name);self.client=self.app.test_client()
        self.client.get('/login')
        with self.client.session_transaction() as s:self.token=s['csrf']
        r=self.client.post('/login',data={'csrf':self.token,'name':'admin','password':'a long passphrase 123'})
        self.assertEqual(r.status_code,302)
        with self.client.session_transaction() as s:self.token=s['csrf']
        self.h={'X-CSRF-Token':self.token}
        w=Workbook();s=w.active;s.title='DU LIEU CO SO';s['E4']='001';s['F4']='Nhân sự thử nghiệm';s['B4']='Kỹ thuật viên';s['N4']='PX1';s['G4']='01/02/1990';s['H4']='=1+1';s.merge_cells('A1:B1')
        w.create_sheet('Báo cáo')['A1']="='DU LIEU CO SO'!F4"
        f=io.BytesIO();w.save(f);f.seek(0)
        r=self.client.post('/api/import',headers=self.h,data={'file':(f,'test.xlsx')})
        self.assertEqual(r.status_code,200);self.bid=r.json['id']
    def tearDown(self):self.tmp.cleanup()
    def test_grid_edit_export_and_conflict(self):
        b='/api/book/'+self.bid
        data=self.client.get(b+'/sheet').json
        self.assertEqual(len(data['rows']),4)
        d={'sheet':'DU LIEU CO SO','cell':'F4','value':'Tên đã sửa','type':'text','revision':0}
        self.assertEqual(self.client.post(b+'/cell',json=d,headers=self.h).status_code,200)
        self.assertEqual(self.client.post(b+'/cell',json=d,headers=self.h).status_code,409)
        d.update(cell='H4',revision=1)
        self.assertEqual(self.client.post(b+'/cell',json=d,headers=self.h).status_code,400)
        d.update(cell='G4',type='date',value='1991-03-04')
        self.assertEqual(self.client.post(b+'/cell',json=d,headers=self.h).status_code,200)
        data=self.client.get(b+'/sheet').json
        self.assertEqual(data['rows'][3]['cells'][6]['value'],'04/03/1991')
        r=self.client.get(b+'/export');w=load_workbook(io.BytesIO(r.data))
        self.assertEqual(w.worksheets[0]['F4'].value,'Tên đã sửa')
        self.assertEqual(w.worksheets[0]['H4'].value,'=1+1')
        self.assertEqual(w.worksheets[0]['E4'].value,'001')
        self.assertIn('A1:B1',[str(m) for m in w.worksheets[0].merged_cells])
    def test_roles_csrf_and_backup(self):
        self.assertEqual(self.client.post('/api/users',json={}).status_code,403)
        self.assertEqual(self.client.post('/api/users',headers=self.h,json={'name':'reader','password':'reader password 123','role':'viewer'}).status_code,200)
        r=self.client.get('/api/backup');z=zipfile.ZipFile(io.BytesIO(r.data));self.assertIn('hr.sqlite',z.namelist());self.assertNotIn('session.key',z.namelist())
        other=self.app.test_client();other.get('/login')
        with other.session_transaction() as s:token=s['csrf']
        other.post('/login',data={'csrf':token,'name':'reader','password':'reader password 123'})
        with other.session_transaction() as s:token=s['csrf']
        self.assertEqual(other.get('/api/books').status_code,200)
        self.assertEqual(other.get('/api/backup').status_code,403)
        self.assertEqual(other.post('/api/book/'+self.bid+'/cell',headers={'X-CSRF-Token':token},json={}).status_code,403)
        anonymous=self.app.test_client();self.assertEqual(anonymous.get('/api/books').status_code,401)
    def test_decisions_and_template_split_runs(self):
        payload={'book':self.bid,'row':4,'kind':'transfer','number':'01/QĐ','effective':'2026-11-01','reason':'Căn cứ thử nghiệm','new_title':'Trưởng ca','new_unit':'PX2','signer':'Người ký thử'}
        for kind in ['transfer','title','termination']:
            payload['kind']=kind;r=self.client.post('/api/decisions',headers=self.h,json=payload)
            self.assertEqual(r.status_code,200)
            doc=Document(io.BytesIO(r.data));self.assertIn('BẢN DỰ THẢO',doc.paragraphs[0].text)
        doc=Document();p=doc.add_paragraph();p.add_run('Người: {{ho_');p.add_run('ten}} — {{ma_nv}}')
        out=io.BytesIO();doc.save(out);out.seek(0)
        self.assertEqual(self.client.post('/api/templates/transfer',headers=self.h,data={'file':(out,'template.docx')}).status_code,200)
        payload['kind']='transfer';r=self.client.post('/api/decisions',headers=self.h,json=payload)
        doc=Document(io.BytesIO(r.data));self.assertEqual(doc.paragraphs[0].text,'Người: Nhân sự thử nghiệm — 001')
        self.assertEqual(self.client.get('/api/book/'+self.bid+'/people').json[0]['unit'],'PX1')
    def test_login_and_pages(self):
        self.assertEqual(self.client.get('/').status_code,200)
        r=self.client.get('/static/app.js');self.assertEqual(r.status_code,200);r.close()
        self.assertEqual(self.client.get('/api/book/'+self.bid).json['external_links'],0)
        self.assertEqual(len(self.client.get('/api/book/'+self.bid+'/people').json),1)
        r=self.client.get('/api/book/'+self.bid+'/sheet?q=khong-tim-thay');self.assertEqual(r.json['rows'],[])

if __name__=='__main__':unittest.main()

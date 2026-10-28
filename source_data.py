"""Read the original personnel sheet without executing Excel formulas."""
import unicodedata
from collections import Counter
from openpyxl.utils import get_column_letter

END_COLUMNS = (197, 213)  # GO: company closing code; HE: TKV closing code
EXCEL_ERRORS = {'#REF!','#N/A','#VALUE!','#NAME?','#DIV/0!','#NUM!','#SPILL!','#NULL!'}

def normalized(value):
    return ''.join(ch for ch in unicodedata.normalize('NFD',str(value or '').casefold()) if unicodedata.category(ch)!='Mn').replace('đ','d').strip()

def value_at(sheet,cached,changes,row,col):
    cell=sheet._cells.get((row,col))
    if cell is None:return changes.get((sheet.title,f'{get_column_letter(col)}{row}'))
    if (sheet.title,cell.coordinate) in changes:return changes[(sheet.title,cell.coordinate)]
    if cell.data_type=='f':
        saved=cached._cells.get((row,col))
        return saved.value if saved is not None else None
    return cell.value

def text_at(sheet,cached,changes,row,col):
    value=value_at(sheet,cached,changes,row,col)
    return str(value).strip() if value is not None else ''

def closing_columns(sheet):
    out=[]
    for col in END_COLUMNS:
        labels=[sheet._cells[(r,col)].value for r in (2,3) if (r,col) in sheet._cells]
        if any('cuoi ki' in normalized(v) or 'cuoi ky' in normalized(v) for v in labels):out.append(col)
    return out

def source_status(sheet,cached,changes,row):
    cols=closing_columns(sheet)
    if not cols:return dict(status='review',end_code='',end_column='',status_reason='Không tìm thấy cột mã cuối kỳ GO/HE có tiêu đề Cuối kỳ.')
    readings=[]
    for col in cols:
        cell=sheet._cells.get((row,col));raw=value_at(sheet,cached,changes,row,col);text=str(raw).strip() if raw is not None else ''
        patched=cell is not None and (sheet.title,cell.coordinate) in changes
        if text in EXCEL_ERRORS or (cell is not None and cell.data_type=='f' and raw is None and not patched):status='review'
        else:status='terminated' if not text or text in ('0','0.0') else 'active'
        readings.append((col,text,status))
    col,text,status=readings[0]
    reason='Có mã cuối kỳ: đang làm.' if status=='active' else 'Không còn mã cuối kỳ: đã nghỉ; vẫn giữ hồ sơ.' if status=='terminated' else 'Mã cuối kỳ bị lỗi hoặc công thức chưa có kết quả lưu.'
    if len({r[2] for r in readings})>1:
        status='review';reason='Mã cuối kỳ GO và HE không thống nhất; cần đối chiếu.'
    return dict(status=status,end_code=text,end_column='GO' if col==197 else 'HE',status_reason=reason)

def education_at(sheet,cached,changes,row):
    flags=[]
    for col,key in [(51,'university'),(78,'college'),(105,'intermediate'),(132,'vocational'),(159,'untrained')]:
        text=text_at(sheet,cached,changes,row,col)
        if text and text not in EXCEL_ERRORS and text not in ('0','0.0'):flags.append(key)
    return flags[0] if len(flags)==1 else 'unknown'

def is_person_row(sheet,cached,changes,row):
    # Do not treat summary formula results as personnel names.
    name=text_at(sheet,cached,changes,row,6)
    if not any(ch.isalpha() for ch in name):return False
    if not text_at(sheet,cached,changes,row,5):
        for col in closing_columns(sheet):
            cell=sheet._cells.get((row,col));saved=value_at(sheet,cached,changes,row,col)
            if cell is not None and cell.data_type=='f' and isinstance(saved,(int,float)):return False
    return normalized(name).rstrip(':') not in ('tong','tong so','tong cong','cong','cong chung')

def source_filters(sheet,cached,changes):
    statuses=Counter();units=set()
    for row in range(4,sheet.max_row+1):
        if not is_person_row(sheet,cached,changes,row):continue
        statuses[source_status(sheet,cached,changes,row)['status']]+=1
        unit=text_at(sheet,cached,changes,row,206) or text_at(sheet,cached,changes,row,14)
        if unit:units.add(unit)
    return dict(status_counts=dict(statuses),units=sorted(units),closing_columns=['GO' if c==197 else 'HE' for c in closing_columns(sheet)])

def matches_source(sheet,cached,changes,row,filters):
    if not is_person_row(sheet,cached,changes,row):return False
    state=source_status(sheet,cached,changes,row)['status']
    if filters.get('status') not in ('',None,'all',state):return False
    unit=text_at(sheet,cached,changes,row,206) or text_at(sheet,cached,changes,row,14)
    if filters.get('unit') and filters['unit']!=unit:return False
    if filters.get('education') and filters['education']!=education_at(sheet,cached,changes,row):return False
    q=filters.get('q','').strip().casefold()
    if q:
        width=getattr(sheet,'_emga_source_width',None)
        if width is None:width=sheet.max_column;sheet._emga_source_width=width
        if not any(q in text_at(sheet,cached,changes,row,col).casefold() for col in range(1,width+1)):return False
    return True
